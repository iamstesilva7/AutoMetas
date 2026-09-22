from flask import Flask, request, jsonify, send_from_directory
from pathlib import Path
from datetime import datetime, date
from io import BytesIO
import traceback
import re
import unicodedata

import pythoncom
import win32com.client
from openpyxl import load_workbook
from openpyxl.styles import Alignment


# ============================================================
# CONFIGURAÇÃO GERAL
# ============================================================

app = Flask(__name__)
BUILD = "DIAS-DINAMICOS-20260829"

BASE_DIR = Path(__file__).resolve().parent
MODELO_EXCEL = BASE_DIR / "METAS.xlsx"
PASTA_SAIDA = BASE_DIR / "metas_geradas"
PASTA_SAIDA.mkdir(parents=True, exist_ok=True)


# ============================================================
# FUNÇÕES GERAIS
# ============================================================

def texto_seguro(valor):
    try:
        return str(valor or "").strip()
    except Exception:
        return ""


def normalizar_texto(valor):
    texto = texto_seguro(valor).lower()
    texto = unicodedata.normalize("NFD", texto)
    texto = "".join(c for c in texto if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", texto).strip()


def normalizar_nome_arquivo(nome):
    nome = texto_seguro(nome)
    nome = re.sub(r'[<>:"/\\|?*]+', "_", nome)
    nome = re.sub(r"\s+", " ", nome).strip().rstrip(".")
    return (nome or "Metas")[:180]


def numero(valor):
    if valor is None or valor == "":
        return 0.0
    if isinstance(valor, (int, float)):
        return float(valor)
    texto = texto_seguro(valor).replace("R$", "").strip()
    if not texto:
        return 0.0
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except (TypeError, ValueError):
        return 0.0


def mes_vigente():
    meses = [
        "Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho",
        "Julho", "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"
    ]
    agora = datetime.now()
    return f"{meses[agora.month - 1]} de {agora.year}"


DIAS_SEMANA_PT = [
    "segunda-feira",
    "terça-feira",
    "quarta-feira",
    "quinta-feira",
    "sexta-feira",
    "sábado",
    "domingo",
]

MESES_PT = [
    "janeiro",
    "fevereiro",
    "março",
    "abril",
    "maio",
    "junho",
    "julho",
    "agosto",
    "setembro",
    "outubro",
    "novembro",
    "dezembro",
]


def formatar_data_extenso_pt_br(data_obj):
    """
    Retorna a data como texto, sem depender da configuração regional do Excel.

    Exemplo:
        sábado, 29 de agosto de 2026.
    """
    return (
        f"{DIAS_SEMANA_PT[data_obj.weekday()]}, "
        f"{data_obj.day} de "
        f"{MESES_PT[data_obj.month - 1]} de "
        f"{data_obj.year}."
    )


# ============================================================
# OUTLOOK
# ============================================================
# Não existe login no AutoMeta.
# O sistema usa o Outlook clássico já aberto no Windows.
#
# IMPORTANTE:
# Em alguns perfis corporativos o Outlook mostra uma caixa nova no painel,
# mas Session.Accounts ainda devolve uma conta antiga/desativada.
# Por isso este código olha TANTO Accounts QUANTO Stores (caixas abertas).
#
# Prioridade:
# 1) caixa @petzcobasi.com.br encontrada nos Stores;
# 2) conta @petzcobasi.com.br encontrada em Accounts;
# 3) CurrentUser;
# 4) primeira conta disponível.
# ============================================================

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\\.[A-Z]{2,}", re.I)
DOMINIO_PRIORITARIO = "@petzcobasi.com.br"


def iniciar_com():
    pythoncom.CoInitialize()


def finalizar_com():
    try:
        pythoncom.CoUninitialize()
    except Exception:
        pass


def obter_outlook():
    iniciar_com()
    erros = []

    try:
        outlook = win32com.client.GetActiveObject("Outlook.Application")
        print("OUTLOOK: conectado à instância já aberta.")
        return outlook
    except Exception as erro:
        erros.append(f"GetActiveObject: {erro}")

    try:
        outlook = win32com.client.Dispatch("Outlook.Application")
        print("OUTLOOK: conectado via Dispatch.")
        return outlook
    except Exception as erro:
        erros.append(f"Dispatch: {erro}")

    finalizar_com()
    raise RuntimeError(
        "Não foi possível acessar o Outlook clássico do Windows. "
        "Confirme se o Outlook desktop clássico está aberto. "
        + " | ".join(erros)
    )


def extrair_email(texto):
    valor = texto_seguro(texto)
    if not valor:
        return ""
    achado = EMAIL_RE.search(valor)
    return achado.group(0) if achado else ""


def smtp_address_entry(address_entry):
    if address_entry is None:
        return ""

    try:
        tipo = texto_seguro(getattr(address_entry, "Type", "")).upper()

        if tipo == "EX":
            try:
                exchange_user = address_entry.GetExchangeUser()
                if exchange_user:
                    smtp = texto_seguro(exchange_user.PrimarySmtpAddress)
                    if smtp:
                        return smtp
            except Exception:
                pass

        endereco = texto_seguro(getattr(address_entry, "Address", ""))
        if "@" in endereco:
            return endereco

        try:
            PR_SMTP_ADDRESS = "http://schemas.microsoft.com/mapi/proptag/0x39FE001E"
            smtp = texto_seguro(
                address_entry.PropertyAccessor.GetProperty(PR_SMTP_ADDRESS)
            )
            if smtp:
                return smtp
        except Exception:
            pass

    except Exception:
        pass

    return ""


def dados_da_conta(conta):
    nome = ""
    smtp = ""

    try:
        nome = texto_seguro(getattr(conta, "DisplayName", ""))
    except Exception:
        pass

    try:
        smtp = texto_seguro(getattr(conta, "SmtpAddress", ""))
    except Exception:
        pass

    return {"display_name": nome, "smtp": smtp}


def usuario_atual_outlook(outlook):
    nome = ""
    email = ""

    try:
        current_user = outlook.Session.CurrentUser
        nome = texto_seguro(getattr(current_user, "Name", ""))
        try:
            email = smtp_address_entry(current_user.AddressEntry)
        except Exception:
            pass
    except Exception:
        pass

    return {"nome": nome, "email": email}


def listar_contas_outlook(outlook):
    resultado = []
    contas = outlook.Session.Accounts

    for indice in range(1, int(contas.Count) + 1):
        conta = contas.Item(indice)
        info = dados_da_conta(conta)
        resultado.append({
            "objeto": conta,
            "nome": info["display_name"],
            "email": info["smtp"],
        })

    return resultado


def listar_stores_outlook(outlook):
    """Lista as caixas que aparecem no painel esquerdo do Outlook."""
    resultado = []
    stores = outlook.Session.Stores

    for indice in range(1, int(stores.Count) + 1):
        store = stores.Item(indice)
        nome = ""
        email = ""

        try:
            nome = texto_seguro(store.DisplayName)
        except Exception:
            pass

        # Muitas caixas corporativas usam o próprio e-mail como DisplayName.
        email = extrair_email(nome)

        # Tenta também o nome da pasta raiz.
        if not email:
            try:
                raiz = store.GetRootFolder()
                nome_raiz = texto_seguro(raiz.Name)
                email = extrair_email(nome_raiz)
                if not nome:
                    nome = nome_raiz
            except Exception:
                pass

        resultado.append({
            "objeto": store,
            "nome": nome,
            "email": email,
        })

    return resultado


def escolher_email_outlook(outlook):
    """
    Descobre o endereço que realmente deve ser usado no AutoMeta.

    Retorna:
        conta_send_using: Account COM correspondente, se existir
        usuario: {nome, email}
    """
    contas = listar_contas_outlook(outlook)
    stores = listar_stores_outlook(outlook)
    atual = usuario_atual_outlook(outlook)

    print(f"OUTLOOK: {len(contas)} conta(s) em Session.Accounts")
    for i, item in enumerate(contas, start=1):
        print(f"  ACCOUNT {i}: {item['nome']} | {item['email']}")

    print(f"OUTLOOK: {len(stores)} Store(s)/caixa(s) aberta(s)")
    for i, item in enumerate(stores, start=1):
        print(f"  STORE {i}: {item['nome']} | {item['email']}")

    # 1) Prioriza a nova caixa corporativa visível no Outlook.
    candidatos = []

    for item in stores:
        email = texto_seguro(item.get("email"))
        if email:
            candidatos.append(("store", item))

    for item in contas:
        email = texto_seguro(item.get("email"))
        if email:
            candidatos.append(("account", item))

    if atual.get("email"):
        candidatos.append(("current", {
            "objeto": None,
            "nome": atual.get("nome", ""),
            "email": atual.get("email", ""),
        }))

    escolhido = None

    # Primeiro procura especificamente o domínio novo da empresa.
    for origem, item in candidatos:
        email = texto_seguro(item.get("email")).lower()
        if email.endswith(DOMINIO_PRIORITARIO):
            escolhido = (origem, item)
            break

    # Se não houver @petzcobasi, usa o CurrentUser se for SMTP válido.
    if escolhido is None and atual.get("email"):
        escolhido = ("current", {
            "objeto": None,
            "nome": atual.get("nome", ""),
            "email": atual.get("email", ""),
        })

    # Depois qualquer endereço válido encontrado.
    if escolhido is None and candidatos:
        escolhido = candidatos[0]

    if escolhido is None:
        raise RuntimeError(
            "O Outlook abriu, mas não consegui encontrar nenhum endereço SMTP "
            "nas contas nem nas caixas abertas."
        )

    origem, item = escolhido
    email_escolhido = texto_seguro(item.get("email"))
    nome_escolhido = texto_seguro(item.get("nome")) or email_escolhido

    # Se houver Account com o mesmo SMTP, usa SendUsingAccount.
    conta_send_using = None
    for conta_info in contas:
        if (
            texto_seguro(conta_info.get("email")).lower()
            == email_escolhido.lower()
        ):
            conta_send_using = conta_info["objeto"]
            break

    print(
        "OUTLOOK: remetente escolhido -> "
        f"{nome_escolhido} | {email_escolhido} | origem={origem}"
    )

    if conta_send_using is None:
        print(
            "OUTLOOK: a caixa escolhida não aparece em Session.Accounts. "
            "O envio usará SentOnBehalfOfName."
        )

    return conta_send_using, {
        "nome": nome_escolhido,
        "email": email_escolhido,
        "origem": origem,
    }


def usuario_da_conta_padrao(outlook):
    # Mantém o nome da função para compatibilidade com o restante do sistema.
    return escolher_email_outlook(outlook)


# ============================================================
# PLANILHA INDIVIDUAL
# ============================================================

def gerar_planilha(meta):
    if not MODELO_EXCEL.exists():
        raise FileNotFoundError(f"O modelo METAS.xlsx não foi encontrado em: {MODELO_EXCEL}")

    codigo = texto_seguro(meta.get("codigo"))
    loja = texto_seguro(meta.get("loja")) or codigo
    metas_loja = meta.get("metasDiarias") or {}
    metas_servico = meta.get("servicosDiarios") or {}

    if not codigo:
        raise ValueError("Código da loja não informado.")
    if not metas_loja:
        raise ValueError(f"A loja {codigo} não possui metas diárias.")

    wb = load_workbook(MODELO_EXCEL)
    if "MODELO LOJA" not in wb.sheetnames:
        raise ValueError("A aba 'MODELO LOJA' não existe no arquivo METAS.xlsx.")

    ws = wb["MODELO LOJA"]
    ws["B3"] = loja

    datas = sorted(
        metas_loja.keys(),
        key=lambda d: datetime.strptime(d, "%d/%m/%Y")
    )

    if len(datas) > 31:
        raise ValueError(f"A loja {codigo} possui mais de 31 dias de meta.")

    # ========================================================
    # DATAS E METAS DIÁRIAS
    # ========================================================
    for indice in range(31):
        linha = 5 + indice

        if indice < len(datas):
            data_texto = datas[indice]
            data_obj = datetime.strptime(data_texto, "%d/%m/%Y")

            # Mantém a data no formato que já estava funcionando.
            ws.cell(linha, 2).value = formatar_data_extenso_pt_br(data_obj)
            ws.cell(linha, 2).number_format = "@"

            ws.cell(linha, 3).value = numero(
                metas_loja.get(data_texto)
            )
            ws.cell(linha, 4).value = numero(
                metas_servico.get(data_texto, 0)
            )
        else:
            ws.cell(linha, 2).value = None
            ws.cell(linha, 3).value = None
            ws.cell(linha, 4).value = None

    # ========================================================
    # REMOVE AS LINHAS QUE SOBRAM NO MÊS
    # ========================================================
    # O modelo comporta até 31 dias (linhas 5 a 35).
    # Se o mês tiver 30, 29 ou 28 dias, as linhas excedentes
    # são removidas fisicamente e a linha TOTAL sobe automaticamente.
    quantidade_dias = len(datas)
    linhas_excedentes = 31 - quantidade_dias

    if linhas_excedentes > 0:
        primeira_linha_excedente = 5 + quantidade_dias
        ws.delete_rows(
            primeira_linha_excedente,
            linhas_excedentes
        )

        print(
            f"MÊS COM {quantidade_dias} DIAS: "
            f"{linhas_excedentes} linha(s) excedente(s) removida(s)."
        )

    # ========================================================
    # TOTAL - encontra automaticamente a linha TOTAL do modelo
    # ========================================================
    # Não dependemos mais de a linha TOTAL ser sempre a 36.
    # Primeiro procura "VALOR TOTAL" ou "TOTAL" no próprio modelo.
    linha_total = None

    for linha_busca in range(1, min(ws.max_row, 80) + 1):
        textos_linha = []

        for coluna_busca in range(1, min(ws.max_column, 20) + 1):
            valor = ws.cell(linha_busca, coluna_busca).value

            if valor is not None:
                textos_linha.append(normalizar_texto(valor))

        texto_linha = " ".join(textos_linha)

        if "valor total" in texto_linha:
            linha_total = linha_busca
            break

    # Se não encontrar "VALOR TOTAL", procura qualquer célula com "TOTAL".
    if linha_total is None:
        for linha_busca in range(1, min(ws.max_row, 80) + 1):
            encontrou_total = False

            for coluna_busca in range(1, min(ws.max_column, 20) + 1):
                valor = normalizar_texto(
                    ws.cell(linha_busca, coluna_busca).value
                )

                if valor == "total" or valor.startswith("total ") or valor.endswith(" total"):
                    linha_total = linha_busca
                    encontrou_total = True
                    break

            if encontrou_total:
                break

    # Compatibilidade com o modelo antigo.
    if linha_total is None:
        linha_total = 36

    # Soma exatamente o que foi gravado na planilha individual.
    ultima_linha_dia = 4 + quantidade_dias

    total_loja = sum(
        numero(ws.cell(linha, 3).value)
        for linha in range(5, ultima_linha_dia + 1)
    )

    total_servico = sum(
        numero(ws.cell(linha, 4).value)
        for linha in range(5, ultima_linha_dia + 1)
    )

    # Grava o número final, não uma fórmula sem cache.
    ws.cell(linha_total, 3).value = total_loja
    ws.cell(linha_total, 4).value = total_servico

    # Formato numérico para o total aparecer normalmente no Excel.
    ws.cell(linha_total, 3).number_format = '#,##0.00'
    ws.cell(linha_total, 4).number_format = '#,##0.00'

    # Fórmulas de realizado/percentual acompanham a linha TOTAL encontrada.
    ws.cell(linha_total, 6).value = (
        f"=SUM(F5:F{ultima_linha_dia})"
    )
    ws.cell(linha_total, 7).value = (
        f"=IFERROR(F{linha_total}/C{linha_total},0)"
    )
    ws.cell(linha_total, 10).value = (
        f"=SUM(J5:J{ultima_linha_dia})"
    )
    ws.cell(linha_total, 11).value = (
        f"=IFERROR(J{linha_total}/D{linha_total},0)"
    )

    print(f"LINHA TOTAL DETECTADA: {linha_total}")
    print(f"TOTAL LOJA {codigo}: {total_loja}")
    print(f"TOTAL PET ANJO {codigo}: {total_servico}")

    # ========================================================
    # CENTRALIZAÇÃO DA PLANILHA
    # ========================================================
    # Código/nome da loja e área de título.
    for linha in range(1, 4):
        for coluna in range(2, 12):  # B até K
            celula = ws.cell(linha, coluna)
            celula.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # Cabeçalhos da tabela.
    for coluna in range(2, 12):  # B até K
        ws.cell(4, coluna).alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Conteúdo da tabela: somente os dias que existem no mês.
    for linha in range(5, ultima_linha_dia + 1):
        for coluna in range(2, 12):  # B até K
            ws.cell(linha, coluna).alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

    # Linha TOTAL encontrada automaticamente.
    for coluna in range(2, 12):
        ws.cell(linha_total, coluna).alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )

    # Mantém o Excel configurado para recalcular as fórmulas restantes.
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcOnSave = True
    wb.calculation.calcCompleted = False
    wb.calculation.calcId = 0

    nome_arquivo = normalizar_nome_arquivo(
        f"Metas_{codigo}_{loja}"
    ) + ".xlsx"

    caminho = PASTA_SAIDA / nome_arquivo
    wb.save(caminho)

    print(f"PLANILHA GERADA: {caminho}")
    print(f"TOTAL LOJA {codigo}: {total_loja}")
    print(f"TOTAL PET ANJO {codigo}: {total_servico}")

    return caminho


# ============================================================
# LEITURA DA PLANILHA PRINCIPAL
# ============================================================

def cabecalho_data(valor):
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")

    texto = texto_seguro(valor)
    if not texto:
        return ""

    padrao = re.match(r"^(\d{1,2})[./-](\d{1,2})[./-](\d{2,4})$", texto)
    if not padrao:
        return ""

    dia, mes, ano = map(int, padrao.groups())
    if ano < 100:
        ano += 2000

    try:
        return datetime(ano, mes, dia).strftime("%d/%m/%Y")
    except ValueError:
        return ""


def valor_json(valor):
    if isinstance(valor, (datetime, date)):
        return valor.strftime("%d/%m/%Y")
    if valor is None:
        return ""
    if isinstance(valor, (str, int, float, bool)):
        return valor
    return str(valor)


def ler_excel_em_objetos(conteudo):
    wb = load_workbook(BytesIO(conteudo), data_only=True, read_only=True)

    melhor_ws = None
    melhor_pontuacao = -1
    melhor_linha_cabecalho = None

    for ws in wb.worksheets:
        limite = min(ws.max_row, 60)
        linhas = list(ws.iter_rows(min_row=1, max_row=limite, values_only=True))

        for indice, linha in enumerate(linhas, start=1):
            normalizados = [normalizar_texto(v) for v in linha]
            possui_filial = any(
                v in {"filial", "codigo filial", "codigo da filial", "loja", "codigo loja", "codigo da loja"}
                for v in normalizados
            )
            quantidade_datas = sum(1 for v in linha if cabecalho_data(v))
            pontuacao = (100 if possui_filial else 0) + quantidade_datas

            if pontuacao > melhor_pontuacao:
                melhor_pontuacao = pontuacao
                melhor_ws = ws
                melhor_linha_cabecalho = indice

    if melhor_ws is None or melhor_linha_cabecalho is None or melhor_pontuacao < 101:
        raise ValueError(
            "Não consegui identificar uma linha de cabeçalho com FILIAL/LOJA e colunas de datas."
        )

    linha_header = next(
        melhor_ws.iter_rows(
            min_row=melhor_linha_cabecalho,
            max_row=melhor_linha_cabecalho,
            values_only=True,
        )
    )

    headers = []
    for indice, valor in enumerate(linha_header, start=1):
        data_formatada = cabecalho_data(valor)
        if data_formatada:
            headers.append(data_formatada)
        else:
            texto = texto_seguro(valor)
            headers.append(texto if texto else f"COL_{indice}")

    dados = []
    for linha in melhor_ws.iter_rows(min_row=melhor_linha_cabecalho + 1, values_only=True):
        if not any(v not in (None, "") for v in linha):
            continue

        obj = {}
        for indice, header in enumerate(headers):
            if indice >= len(linha):
                break
            obj[header] = valor_json(linha[indice])
        dados.append(obj)

    return melhor_ws.title, melhor_linha_cabecalho, dados


# ============================================================
# DIAGNÓSTICO DE REQUISIÇÕES
# ============================================================

@app.before_request
def log_requisicao():
    print(f"[BUILD {BUILD}] {request.method} {request.path}", flush=True)


# ============================================================
# ROTAS
# ============================================================

@app.route("/")
def index():
    response = send_from_directory(str(BASE_DIR), "index1.html")
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return response


@app.route("/app.js")
def app_js():
    response = send_from_directory(str(BASE_DIR), "app.js", mimetype="application/javascript")
    response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
    return response


@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({
        "success": True,
        "build": BUILD,
        "message": "AutoMeta Flask funcionando.",
        "mes_vigente": mes_vigente(),
        "modelo_excel": str(MODELO_EXCEL),
        "modelo_existe": MODELO_EXCEL.exists(),
    })


@app.route("/api/usuario-outlook", methods=["GET"])
def usuario_outlook():
    try:
        print("\n======================================")
        print("IDENTIFICAÇÃO AUTOMÁTICA DO OUTLOOK")
        print("======================================")

        outlook = obter_outlook()
        conta, usuario = usuario_da_conta_padrao(outlook)

        print(f"USUÁRIO OUTLOOK: {usuario['nome']}")
        print(f"E-MAIL OUTLOOK: {usuario['email']}")
        print("======================================\n")

        return jsonify({
            "success": True,
            "usuario": usuario,
            "conta": dados_da_conta(conta) if conta is not None else {
                "display_name": usuario["nome"],
                "smtp": usuario["email"],
            },
        })
    except Exception as erro:
        print("\nERRO AO IDENTIFICAR OUTLOOK")
        traceback.print_exc()
        return jsonify({"success": False, "message": str(erro)}), 500
    finally:
        finalizar_com()


@app.route("/api/outlook-contas", methods=["GET"])
def outlook_contas():
    try:
        outlook = obter_outlook()
        contas = outlook.Session.Accounts
        lista = [dados_da_conta(contas.Item(i)) for i in range(1, contas.Count + 1)]
        return jsonify({"success": True, "contas": lista})
    except Exception as erro:
        print("ERRO /api/outlook-contas:")
        traceback.print_exc()
        return jsonify({"success": False, "message": str(erro)}), 500
    finally:
        finalizar_com()


@app.route("/api/outlook-diagnostico", methods=["GET", "POST"])
def outlook_diagnostico():
    try:
        outlook = obter_outlook()
        conta, usuario = usuario_da_conta_padrao(outlook)
        return jsonify({
            "success": True,
            "usuario": usuario,
            "conta_encontrada": dados_da_conta(conta) if conta is not None else {
                "display_name": usuario["nome"],
                "smtp": usuario["email"],
            },
            "mensagem": "Outlook identificado com sucesso.",
        })
    except Exception as erro:
        print("\nERRO NO DIAGNÓSTICO DO OUTLOOK:")
        traceback.print_exc()
        return jsonify({"success": False, "message": str(erro)}), 500
    finally:
        finalizar_com()


@app.route("/api/ler-planilha", methods=["POST"])
def ler_planilha():
    try:
        arquivo = request.files.get("arquivo")
        if arquivo is None or not arquivo.filename:
            raise ValueError("Nenhuma planilha foi recebida pelo servidor.")

        nome = arquivo.filename.lower()
        if not (nome.endswith(".xlsx") or nome.endswith(".xlsm")):
            raise ValueError("Envie uma planilha .xlsx ou .xlsm.")

        conteudo = arquivo.read()
        if not conteudo:
            raise ValueError("A planilha recebida está vazia.")

        aba, linha_cabecalho, dados = ler_excel_em_objetos(conteudo)

        print("\n======================================")
        print("PLANILHA RECEBIDA PELO FLASK")
        print(f"Arquivo: {arquivo.filename}")
        print(f"Aba: {aba}")
        print(f"Cabeçalho: linha {linha_cabecalho}")
        print(f"Linhas retornadas: {len(dados)}")
        print("======================================\n")

        return jsonify({
            "success": True,
            "aba": aba,
            "linha_cabecalho": linha_cabecalho,
            "dados": dados,
        })
    except Exception as erro:
        print("\nERRO AO LER PLANILHA:")
        traceback.print_exc()
        return jsonify({"success": False, "message": str(erro)}), 500


@app.route("/api/enviar-meta", methods=["POST"])
def enviar_meta():
    destinatario = ""
    caminho_planilha = None
    usuario_outlook = {"nome": "", "email": ""}

    try:
        dados = request.get_json(silent=True) or {}

        destinatario = texto_seguro(dados.get("destinatario"))
        assunto = texto_seguro(dados.get("assunto"))
        corpo = str(dados.get("corpo") or "")
        loja = texto_seguro(dados.get("loja"))
        codigo = texto_seguro(dados.get("codigo"))

        print("\n======================================")
        print("SOLICITAÇÃO DE ENVIO")
        print(f"Loja: {loja}")
        print(f"Código: {codigo}")
        print(f"Destinatário: {destinatario}")

        if not destinatario:
            raise ValueError(f"A loja {loja or codigo} não possui destinatário.")
        if not assunto:
            raise ValueError("Assunto do e-mail não informado.")
        if not corpo:
            raise ValueError("Corpo do e-mail não informado.")

        caminho_planilha = gerar_planilha(dados)

        outlook = obter_outlook()
        conta_remetente, usuario_outlook = usuario_da_conta_padrao(outlook)

        mensagem = outlook.CreateItem(0)
        mensagem.To = destinatario
        mensagem.Subject = assunto
        mensagem.Body = corpo

        if conta_remetente is not None:
            try:
                mensagem.SendUsingAccount = conta_remetente
            except Exception:
                # Fallback comum para COM late binding
                mensagem._oleobj_.Invoke(*(64209, 0, 8, 0, conta_remetente))
        else:
            # A caixa nova aparece no Outlook como Store, mas não como Account.
            # Nesse cenário o Outlook/Exchange deve resolver o endereço da própria
            # caixa e enviar usando a permissão disponível para esse mailbox.
            mensagem.SentOnBehalfOfName = usuario_outlook["email"]

        mensagem.Attachments.Add(str(caminho_planilha))

        print(f"REMETENTE: {usuario_outlook['email'] or usuario_outlook['nome']}")
        print(f"ANEXO: {caminho_planilha}")
        print("ENVIANDO PELO OUTLOOK...")

        mensagem.Send()

        print("E-MAIL ENVIADO COM SUCESSO")
        print("======================================\n")

        return jsonify({
            "success": True,
            "message": f"E-mail da loja {loja or codigo} enviado com sucesso.",
            "usuario": usuario_outlook["nome"],
            "remetente": usuario_outlook["email"],
            "destinatario": destinatario,
            "anexo": str(caminho_planilha),
            "arquivo": caminho_planilha.name,
        })

    except Exception as erro:
        print("\n!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!")
        print("ERRO AO ENVIAR E-MAIL")
        print(str(erro))
        traceback.print_exc()
        print("!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!!\n")

        return jsonify({
            "success": False,
            "message": str(erro),
            "usuario": usuario_outlook.get("nome", ""),
            "remetente": usuario_outlook.get("email", ""),
            "destinatario": destinatario,
            "anexo_gerado": str(caminho_planilha) if caminho_planilha else "",
        }), 500
    finally:
        finalizar_com()


# ============================================================
# INICIAR SERVIDOR
# ============================================================

if __name__ == "__main__":
    print("\n======================================")
    print("AUTOMETA - SERVIDOR FLASK")
    print(f"BUILD: {BUILD}")
    print("======================================")
    print("Servidor: http://127.0.0.1:5055")
    print(f"Mês vigente: {mes_vigente()}")
    print(f"Modelo: {MODELO_EXCEL}")
    print(f"Modelo encontrado: {MODELO_EXCEL.exists()}")
    print("Usuário: identificado automaticamente pelo Outlook")
    print("Remetente: conta padrão do Outlook desta máquina")
    print("Anexo: cópia individual de METAS.xlsx")
    print("======================================\n")

    app.run(
        host="127.0.0.1",
        port=5055,
        debug=False,
        use_reloader=False,
    )
